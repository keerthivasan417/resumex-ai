"""Deterministic text extraction and section detection for resumes."""

from dataclasses import dataclass
from io import BytesIO
import re
from zipfile import BadZipFile, ZipFile

import fitz
from docx import Document
from docx.opc.exceptions import PackageNotFoundError


class ExtractionError(ValueError):
    """Raised when a supported resume file cannot be read."""


@dataclass(frozen=True)
class DetectedSection:
    """A contiguous, labelled region of raw resume text."""

    section_type: str
    title: str | None
    position: int
    content: str


_SECTION_HEADINGS: dict[str, set[str]] = {
    "contact": {"contact", "contact information", "contact details"},
    "summary": {"summary", "professional summary", "profile", "objective", "career objective"},
    "skills": {"skills", "technical skills", "core competencies", "technologies", "skills and tools"},
    "experience": {"experience", "work experience", "professional experience", "employment history", "work history"},
    "education": {"education", "academic background", "qualifications"},
    "projects": {"projects", "personal projects", "selected projects", "academic projects"},
    "certifications": {"certifications", "certificates", "licenses and certifications"},
    "links": {"links", "online presence", "profiles", "portfolio"},
}


def extract_text(file_bytes: bytes, extension: str) -> str:
    """Extract text from a validated PDF or DOCX byte stream."""
    if extension == ".pdf":
        return _extract_pdf_text(file_bytes)
    if extension == ".docx":
        return _extract_docx_text(file_bytes)
    raise ExtractionError("Unsupported file extension.")


def _extract_pdf_text(file_bytes: bytes) -> str:
    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as document:
            return "\n".join(page.get_text("text") for page in document).strip()
    except (fitz.FileDataError, RuntimeError, ValueError) as error:
        raise ExtractionError("The PDF could not be read.") from error


def _extract_docx_text(file_bytes: bytes) -> str:
    try:
        with ZipFile(BytesIO(file_bytes)) as archive:
            if "word/document.xml" not in archive.namelist():
                raise ExtractionError("The DOCX file is missing its document content.")
        document = Document(BytesIO(file_bytes))
    except (BadZipFile, KeyError, PackageNotFoundError, ValueError) as error:
        raise ExtractionError("The DOCX file could not be read.") from error

    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines).strip()


def detect_sections(raw_text: str) -> list[DetectedSection]:
    """Split text into stable, heading-based resume sections without inference."""
    lines = [line.strip() for line in raw_text.splitlines()]
    sections: list[DetectedSection] = []
    current_type = "other"
    current_title: str | None = None
    current_lines: list[str] = []

    for line in lines:
        heading_type = _section_type_for_heading(line)
        if heading_type is not None:
            _append_section(sections, current_type, current_title, current_lines)
            current_type = heading_type
            current_title = line
            current_lines = []
        else:
            current_lines.append(line)

    _append_section(sections, current_type, current_title, current_lines)

    if not sections and raw_text.strip():
        return [DetectedSection("other", None, 0, raw_text.strip())]

    return [
        DetectedSection(section.section_type, section.title, position, section.content)
        for position, section in enumerate(sections)
    ]


def _section_type_for_heading(line: str) -> str | None:
    normalized = re.sub(r"[^a-z0-9]+", " ", line.casefold()).strip()
    for section_type, headings in _SECTION_HEADINGS.items():
        if normalized in headings:
            return section_type
    return None


def _append_section(
    sections: list[DetectedSection],
    section_type: str,
    title: str | None,
    lines: list[str],
) -> None:
    content = "\n".join(line for line in lines if line).strip()
    if not content:
        return
    if title is None and section_type == "other" and _looks_like_contact(content):
        section_type = "contact"
    sections.append(DetectedSection(section_type, title, 0, content))


def _looks_like_contact(content: str) -> bool:
    first_lines = content.splitlines()[:4]
    first_block = " ".join(first_lines)
    return "@" in first_block or "http://" in first_block or "https://" in first_block
