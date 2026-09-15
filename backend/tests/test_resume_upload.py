"""Focused tests for deterministic resume upload processing."""

from io import BytesIO
from uuid import uuid4

import fitz
import pytest
from docx import Document

from app.services.resume_extraction import detect_sections, extract_text
from app.services.resume_upload import UploadValidationError, persist_resume, validate_upload


def test_pdf_text_extraction() -> None:
    """PyMuPDF extracts text from a valid PDF byte stream."""
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Jane Doe\nSkills\nPython")
    pdf_bytes = document.tobytes()
    document.close()

    assert extract_text(pdf_bytes, ".pdf") == "Jane Doe\nSkills\nPython"


def test_docx_text_extraction() -> None:
    """python-docx extracts paragraphs and table content from a DOCX stream."""
    document = Document()
    document.add_paragraph("Jane Doe")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python"
    output = BytesIO()
    document.save(output)

    assert extract_text(output.getvalue(), ".docx") == "Jane Doe\nSkills | Python"


def test_invalid_upload_is_rejected_before_storage() -> None:
    """Unsupported extensions and mismatched PDF signatures are rejected."""
    with pytest.raises(UploadValidationError, match="Only PDF and DOCX"):
        validate_upload("resume.txt", "text/plain", b"plain text")

    with pytest.raises(UploadValidationError, match="invalid file signature"):
        validate_upload("resume.pdf", "application/pdf", b"not a PDF")


def test_section_detection_is_heading_based_and_stable() -> None:
    """Known headings create ordered sections without semantic inference."""
    raw_text = """Jane Doe
jane@example.com
SUMMARY
Backend engineer
SKILLS
Python, FastAPI
EXPERIENCE
Example Corp
EDUCATION
Example University
"""

    sections = detect_sections(raw_text)

    assert [(section.section_type, section.position) for section in sections] == [
        ("contact", 0),
        ("summary", 1),
        ("skills", 2),
        ("experience", 3),
        ("education", 4),
    ]
    assert sections[2].content == "Python, FastAPI"


class _FakeSession:
    """Minimal session double for testing model construction without PostgreSQL."""

    def __init__(self) -> None:
        self.added: list[object] = []
        self.flushed = False

    def add(self, instance: object) -> None:
        self.added.append(instance)

    def flush(self) -> None:
        self.flushed = True


def test_persistence_builds_resume_and_sections() -> None:
    """Persistence assigns raw text and section rows to a single resume aggregate."""
    sections = detect_sections("CONTACT\njane@example.com\nSKILLS\nPython")
    db = _FakeSession()

    resume = persist_resume(
        db,  # type: ignore[arg-type]
        uuid4(),
        "resume.pdf",
        "application/pdf",
        "generated.pdf",
        "CONTACT\njane@example.com\nSKILLS\nPython",
        sections,
    )

    assert db.added == [resume]
    assert db.flushed is True
    assert resume.raw_text == "CONTACT\njane@example.com\nSKILLS\nPython"
    assert [section.section_type for section in resume.sections] == ["contact", "skills"]
